from types import SimpleNamespace

from app.models import Chunk, Document
from app.services import AIService


class FakeVectors(list):
    def tolist(self):
        return list(self)


class FakeEmbedder:
    @staticmethod
    def _vector(text):
        return [1.0, 0.0] if "RAG" in text or "檢索" in text else [0.0, 1.0]

    def encode_document(self, texts, **_kwargs):
        return FakeVectors([self._vector(text) for text in texts])

    def encode_query(self, texts, **_kwargs):
        return FakeVectors([self._vector(text) for text in texts])

    encode = encode_document


class FakeOllama:
    def chat(self, **_kwargs):
        return SimpleNamespace(message=SimpleNamespace(content="RAG 會先檢索教材，再根據片段回答。（第 2 頁）"))


from app.workflows import build_deck_graph, build_qa_graph, build_quiz_graph


def ollama_service():
    service = AIService.__new__(AIService)
    service.provider = "ollama"
    service.model = "test-model"
    service.embedding_model = "test-embedding"
    service.openai = None
    service.gemini_client = None
    service.gemini_api_key = ""
    service.ollama = FakeOllama()
    service._embedder = FakeEmbedder()
    service.qa_graph = build_qa_graph()
    service.deck_graph = build_deck_graph()
    service.quiz_graph = build_quiz_graph()
    return service


def document():
    return Document(
        id="doc",
        name="教材.pdf",
        pages=2,
        chunks=[
            Chunk("一般課程介紹", 1, 0),
            Chunk("RAG 先檢索可信文件，再交給模型生成答案。", 2, 1),
        ],
        size_bytes=100,
    )


def test_huggingface_embeddings_drive_retrieval():
    service = ollama_service()
    doc = document()
    service.index(doc)

    matches = service.retrieve(doc, "什麼是 RAG？")

    assert doc.vectors == [[0.0, 1.0], [1.0, 0.0]]
    assert matches[0][0].page == 2


def test_ollama_generates_answer_with_sources():
    service = ollama_service()
    doc = document()
    service.index(doc)

    answer, sources, mode = service.ask(doc, "什麼是 RAG？")

    assert "第 2 頁" in answer
    assert sources[0].page == 2
    assert mode == "ollama"


def test_ollama_deck_uses_structured_model_output():
    service = ollama_service()
    service._structured_response = lambda _system, _prompt, _schema: {
        "title": "RAG 入門",
        "subtitle": "高中生｜20 分鐘",
        "slides": [
            {
                "title": f"單元 {index}",
                "bullets": ["重點一", "重點二"],
                "speaker_notes": "逐頁講稿",
                "source_pages": [2],
            }
            for index in range(4)
        ],
    }

    deck = service.generate_deck(document(), "高中生", "活潑", 4, 20)

    assert len(deck.slides) == 4
    assert deck.mode == "ollama"
    assert all(slide.speaker_notes for slide in deck.slides)


def test_set_provider_gemini():
    service = ollama_service()
    info = service.set_provider("gemini")
    assert info["provider"] == "gemini"
    assert info["provider_label"] == "Gemini 雲端 API"
    assert info["generation_model"] == "gemini-3.6-flash"
    assert info["embedding_model"] == "gemini-embedding-2"


def test_vision_page_to_markdown():
    service = ollama_service()
    # 測試即使傳入無效影像，亦能安全傳回空字串不遺失程序
    res = service._vision_page_to_markdown(b"fake-image-bytes")
    assert isinstance(res, str)



def test_parse_json_response_with_think_tags():
    from app.services import _parse_json_response
    raw = "<think>思考中... 需要規劃簡報內容。</think>\n```json\n{\"title\": \"測試標題\", \"slides\": []}\n```"
    res = _parse_json_response(raw)
    assert res["title"] == "測試標題"

    # 測試未閉合的 think 標籤
    raw_unclosed = "<think>思考中未閉合\n{\"title\": \"未閉合測試\"}"
    res2 = _parse_json_response(raw_unclosed)
    assert res2["title"] == "未閉合測試"


def test_structured_response_fallback():
    service = ollama_service()
    calls = []
    def fake_chat(model, messages, format, stream, options):
        calls.append(format)
        if isinstance(format, dict):
            raise Exception("400 Bad Request: dict format not supported")
        return SimpleNamespace(message=SimpleNamespace(content='{"title": "備援成功"}'))

    service.ollama.chat = fake_chat
    res = service._structured_response("system", "prompt", {"type": "object"})

    assert len(calls) == 2
    assert calls[0] == {"type": "object"}
    assert calls[1] == "json"
    assert res["title"] == "備援成功"


def test_parse_pdf_multimodal_parallel(monkeypatch):
    import os
    from app.services import parse_pdf

    monkeypatch.setenv("ENABLE_MULTIMODAL_PARSING", "true")
    monkeypatch.setenv("MULTIMODAL_MAX_WORKERS", "4")

    # 模擬 PyMuPDF Document 與 Pages
    class FakePage:
        def __init__(self, idx):
            self.idx = idx
        def get_pixmap(self, dpi=200):
            class FakePix:
                def tobytes(self, fmt):
                    return f"img-{self.idx}".encode()
            pix = FakePix()
            pix.idx = self.idx
            return pix
        def get_text(self):
            return f"Page {self.idx} text"

    class FakePyMuPDFDoc:
        def __init__(self, count=5):
            self.pages = [FakePage(i) for i in range(1, count + 1)]
        def __len__(self):
            return len(self.pages)
        def __iter__(self):
            return iter(self.pages)

    # 模擬 PyMuPDF open
    fake_pymupdf = SimpleNamespace(open=lambda stream, filetype: FakePyMuPDFDoc(5))
    monkeypatch.setitem(__import__("sys").modules, "pymupdf", fake_pymupdf)

    service = ollama_service()
    service._vision_page_to_markdown = lambda img_bytes: f"Vision Markdown content for {img_bytes.decode()}"

    doc = parse_pdf(b"%PDF-test", "test_multipage.pdf", ai_service=service)

    assert doc.pages == 5
    assert len(doc.chunks) == 5
    for i, chunk in enumerate(doc.chunks, 1):
        assert chunk.page == i
        assert f"img-{i}" in chunk.text


def test_prepare_openai_strict_schema():
    from app.services import _prepare_openai_strict_schema
    schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "slides": {
                "type": "array",
                "minItems": 5,
                "items": {
                    "type": "object",
                    "properties": {
                        "name": {"type": "string"},
                    },
                },
            },
        },
    }
    cleaned = _prepare_openai_strict_schema(schema)
    assert cleaned["additionalProperties"] is False
    assert cleaned["required"] == ["title", "slides"]
    assert "minItems" not in cleaned["properties"]["slides"]
    assert cleaned["properties"]["slides"]["items"]["additionalProperties"] is False


def test_openai_text_and_structured_response():
    service = ollama_service()
    service.provider = "openai"
    service.model = "gpt-4o-mini"
    
    chat_calls = []

    class FakeChatCompletions:
        def create(self, model, messages, **kwargs):
            chat_calls.append((model, messages, kwargs))
            if "response_format" in kwargs:
                return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='{"title": "OpenAI 簡報"}'))])
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content="OpenAI 文字回應"))])

    class FakeOpenAI:
        chat = SimpleNamespace(completions=FakeChatCompletions())

    service.openai = FakeOpenAI()

    # 測試 _text_response
    text_res = service._text_response("system", "prompt")
    assert text_res == "OpenAI 文字回應"

    # 測試 _structured_response
    struct_res = service._structured_response("system", "prompt", {"type": "object", "properties": {"title": {"type": "string"}}})
    assert struct_res["title"] == "OpenAI 簡報"
    assert len(chat_calls) == 2


def test_page_needs_vision():
    from app.services import _page_needs_vision

    class PageNoImg:
        def get_images(self):
            return []

    class PageWithImg:
        def get_images(self):
            return [("img1",)]

    # 1. 純長文字 -> 不需要 Vision
    pure_text = "這是一段非常標準的中文教學文字，內容包含歷史與社會學概念介紹，完全沒有任何複雜公式與圖片。" * 2
    assert _page_needs_vision(PageNoImg(), pure_text) is False

    # 2. 含有圖片 -> 需要 Vision
    assert _page_needs_vision(PageWithImg(), pure_text) is True

    # 3. 含有 LaTeX/數學符號 -> 需要 Vision
    math_text = "請計算以下幾何公式：當邊長為 a 時，斜邊長為 \\sqrt{a^2 + b^2} 並且滿足勾股定理。" * 2
    assert _page_needs_vision(PageNoImg(), math_text) is True

    unicode_math = "當溫度升高 ±5℃ 時，能量變化為 E = hν。" * 2
    assert _page_needs_vision(PageNoImg(), unicode_math) is True

    # 4. 含有表格結構 -> 需要 Vision
    table_text = "| 項目 | 數量 |\n| --- | --- |\n| 蘋果 | 10 |" * 2
    assert _page_needs_vision(PageNoImg(), table_text) is True

    # 5. 文字過少 (圖片/掃描頁) -> 需要 Vision
    short_text = "頁碼 1"
    assert _page_needs_vision(PageNoImg(), short_text) is True


def test_parse_pdf_smart_routing(monkeypatch):
    from app.services import parse_pdf

    monkeypatch.setenv("ENABLE_MULTIMODAL_PARSING", "true")
    monkeypatch.setenv("ENABLE_SMART_ROUTING", "true")

    vision_calls = []

    class FakePage1:
        def get_images(self): return []
        def get_pixmap(self, dpi=200): return SimpleNamespace(tobytes=lambda fmt: b"img1")
        def get_text(self): return "這是一段很長很長的純文字教學內容，專門用來測試智慧分流純文字管道。歷史課本第一章節介紹。" * 2

    class FakePage2:
        def get_images(self): return []
        def get_pixmap(self, dpi=200): return SimpleNamespace(tobytes=lambda fmt: b"img2")
        def get_text(self): return "數學第二章：請計算微分方程 \\int_0^\\infty e^{-x} dx 的極限值。" * 2

    class FakePyMuPDFDoc:
        def __init__(self):
            self.pages = [FakePage1(), FakePage2()]
        def __len__(self): return len(self.pages)
        def __iter__(self): return iter(self.pages)

    fake_pymupdf = SimpleNamespace(open=lambda stream, filetype: FakePyMuPDFDoc())
    monkeypatch.setitem(__import__("sys").modules, "pymupdf", fake_pymupdf)

    service = ollama_service()

    def fake_vision(img_bytes):
        vision_calls.append(img_bytes)
        return f"Vision result for {img_bytes.decode()}"

    service._vision_page_to_markdown = fake_vision

    doc = parse_pdf(b"%PDF-test", "test_smart_routing.pdf", ai_service=service)

    assert doc.pages == 2
    assert len(vision_calls) == 1
    assert vision_calls[0] == b"img2"
    assert "歷史課本第一章節" in doc.chunks[0].text
    assert "Vision result for img2" in doc.chunks[1].text


def test_make_pptx_with_visual_diagram():
    from app.models import Deck, Slide
    from app.services import make_pptx
    import io
    from pptx import Presentation

    slide = Slide(
        title="**光電效應機制**與 $E=mc^2$",
        bullets=[
            "**光子入射**：單一光子入射金屬表面",
            "**功函數**：克服功函數 $W$ 逸出電子",
            "**動能守恆**：剩餘能量轉為最大動能 $K_{max} = hf - W$"
        ],
        speaker_notes="各位好，我們來看光電效應的核心推導流程：$K_{max} = hf - W$。",
        source_pages=[1, 2],
        icon="⚡",
        visual_description="光子能量轉換示意圖",
        visual_diagram={
            "diagram_type": "flowchart",
            "steps": [
                {"label": "① 光子入射", "text": "**能量傳遞**：單一光子將能量 $hf$ 傳遞給金屬電子"},
                {"label": "② 克服功函數", "text": "**電子逸出**：克服表面束縛功函數 $W$"},
                {"label": "③ 剩餘動能", "text": "**最大動能**：獲得動能 $K_{max}$"},
            ],
            "takeaway": "核心結論：**光電子動能**僅取決於入射光頻率 $f$，與光強度無關。",
        },
    )
    deck = Deck(
        id="test-deck",
        document_id="doc-1",
        title="量子力學入門",
        subtitle="大學生｜30 分鐘",
        slides=[slide],
        duration=30,
        mode="gemini",
    )
    pptx_bytes = make_pptx(deck)
    assert isinstance(pptx_bytes, bytes)
    assert len(pptx_bytes) > 1000

    # Verify Presentation structure
    prs = Presentation(io.BytesIO(pptx_bytes))
    assert len(prs.slides) == 1
    s0 = prs.slides[0]
    # Shapes should include background, accent, slide number, title, bullets body, card container, card title, 3 step boxes, 1 takeaway box = 11 shapes
    assert len(s0.shapes) >= 10


def test_generate_quiz_service_and_markdown():
    from app.services import make_quiz_markdown

    service = ollama_service()
    doc = document()

    def fake_structured_response(system, prompt, schema):
        if "出題大綱" in system or "出題考點方向" in prompt:
            return {
                "title": "RAG 與向量檢索評量",
                "description": "檢測對 RAG 架構與相似度檢索的掌握度",
                "focal_topics": ["向量相似度計算", "幻覺校驗機制"],
            }
        return {
            "title": "RAG 與向量檢索評量",
            "description": "檢測對 RAG 架構與相似度檢索的掌握度",
            "questions": [
                {
                    "type": "single_choice",
                    "question": "在 RAG 系統中，何者用於衡量文字語意相似度？",
                    "options": ["A. 餘弦相似度 (Cosine Similarity)", "B. 歐式距離的倒數", "C. 字符長度比較", "D. 雜湊值碰撞率"],
                    "answer": "A",
                    "explanation": "餘弦相似度常用於衡量高維向量空間中語意方向的一致性。",
                    "source_pages": [1, 2],
                    "difficulty": "easy",
                },
                {
                    "type": "single_choice",
                    "question": "在 RAG 系統中，何者用於幻覺校驗？",
                    "options": ["A. 內容比對", "B. 隨機猜測", "C. 長度統計", "D. 忽略檢索"],
                    "answer": "A",
                    "explanation": "幻覺校驗比對生成答案與檢索段落。",
                    "source_pages": [1, 2],
                    "difficulty": "easy",
                },
            ],
        }

    service._structured_response = fake_structured_response

    quiz = service.generate_quiz(doc, question_count=2, difficulty="easy")
    assert quiz.title == "RAG 與向量檢索評量"
    assert len(quiz.questions) == 2
    assert quiz.questions[0].answer == "A"
    assert quiz.questions[0].source_pages == [1, 2]

    # 驗證學生版與教師版 Markdown
    student_md = make_quiz_markdown(quiz, teacher_mode=False)
    assert "作答區 / 演算草稿" in student_md
    assert "【標準答案】" not in student_md

    teacher_md = make_quiz_markdown(quiz, teacher_mode=True)
    assert "【標準答案】" in teacher_md
    assert "餘弦相似度常用於衡量" in teacher_md
    assert "教材出處頁碼：第 1, 2 頁" in teacher_md


def test_document_store_disk_persistence(tmp_path):
    from app.models import Chunk, Deck, Document, Handout, HandoutSection, Slide
    from app.services import DocumentStore

    store_dir = str(tmp_path / "custom_store")
    store = DocumentStore(base_dir=store_dir)

    doc = Document(
        id="doc_test_123",
        name="測試教材.pdf",
        pages=5,
        chunks=[Chunk(text="區塊內文 1", page=1, index=0)],
        size_bytes=1024,
    )
    store.add(doc)

    handout = Handout(
        id="handout_test_123",
        document_id="doc_test_123",
        title="持久化測試講義",
        subtitle="測試副標",
        overview="課程導讀",
        sections=[
            HandoutSection(
                title="第一章",
                summary="摘要",
                key_points=["重點一"],
                discussion_questions=["問題一"],
                source_pages=[1],
            )
        ],
        key_takeaways=["核心總結"],
    )
    store.handouts[handout.id] = handout

    # 驗證新建立的 store instance 是否能自動從磁碟載入
    reloaded_store = DocumentStore(base_dir=store_dir)
    assert "doc_test_123" in reloaded_store.documents
    assert reloaded_store.get("doc_test_123").name == "測試教材.pdf"
    assert "handout_test_123" in reloaded_store.handouts
    assert reloaded_store.handouts["handout_test_123"].title == "持久化測試講義"


def test_make_handout_html_katex():
    from app.models import Handout, HandoutSection
    from app.services import make_handout_html

    handout = Handout(
        id="h_1",
        document_id="d_1",
        title="物理講義：$\\vec{F}=m\\vec{a}$",
        subtitle="**公式推導**手冊",
        overview="牛頓第二定律與 **質能等價** $E=mc^2$",
        sections=[
            HandoutSection(
                title="萬有引力 $F=G\\frac{m_1 m_2}{r^2}$",
                summary="重力常數 **$G$** 與 `常數定義`",
                key_points=["加速度 $\\vec{a}$", "**重要結論**：$F_{net} > 0$"],
                discussion_questions=["如何測量 $G$？"],
                source_pages=[1, 2],
            )
        ],
        key_takeaways=["能量守恆定律"],
    )

    html_out = make_handout_html(handout)
    assert "katex.min.css" in html_out
    assert "renderMathInElement" in html_out
    assert "物理講義" in html_out
    assert "<strong>公式推導</strong>" in html_out
    assert "<code>常數定義</code>" in html_out
    assert "$\\vec{F}=m\\vec{a}$" in html_out
    assert "A4 portrait" in html_out


def test_make_handout_docx():
    from app.models import Handout, HandoutSection
    from app.services import make_handout_docx

    handout = Handout(
        id="h_docx_1",
        document_id="d_1",
        title="牛頓運動定律與萬有引力",
        subtitle="高中物理第一單元講義",
        overview="本單元探討物體運動狀態與力的作用關係。",
        sections=[
            HandoutSection(
                title="牛頓第一運動定律",
                summary="慣性定律：合力為零時維持等速或靜止。",
                key_points=["慣性與質量的關係", "日常生活慣性現象"],
                discussion_questions=["為什麼急煞車時人會往前傾？"],
                source_pages=[1, 2],
            )
        ],
        key_takeaways=["力是改變運動狀態的原因"],
    )

    docx_bytes = make_handout_docx(handout)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000  # Valid docx zip file


def test_make_quiz_docx_and_html():
    from app.models import QuizSheet, QuizQuestion
    from app.services import make_quiz_docx, make_quiz_html

    sheet = QuizSheet(
        id="q_docx_1",
        document_id="d_1",
        title="牛頓運動定律單元測驗卷",
        description="檢驗牛頓三大運動定律理解程度",
        duration_minutes=20,
        questions=[
            QuizQuestion(
                id="q1",
                type="single_choice",
                question="一物體質量 2kg，受 10N 合力作用，其加速度為？",
                options=["(A) 2 m/s²", "(B) 5 m/s²", "(C) 10 m/s²", "(D) 20 m/s²"],
                answer="(B)",
                explanation="由 F = ma 可得 a = F/m = 10/2 = 5 m/s²",
                source_pages=[3],
                difficulty="easy",
            )
        ],
    )

    # 1. Test student docx
    student_docx = make_quiz_docx(sheet, teacher_mode=False)
    assert isinstance(student_docx, bytes)
    assert len(student_docx) > 1000

    # 2. Test teacher docx
    teacher_docx = make_quiz_docx(sheet, teacher_mode=True)
    assert isinstance(teacher_docx, bytes)
    assert len(teacher_docx) > 1000

    # 3. Test student HTML
    student_html = make_quiz_html(sheet, teacher_mode=False)
    assert "學生練習測驗卷" in student_html
    assert "作答區" in student_html
    assert "【標準答案】" not in student_html
    assert "katex.min.css" in student_html

    # 4. Test teacher HTML
    teacher_html = make_quiz_html(sheet, teacher_mode=True)
    assert "教師詳解卷" in teacher_html
    assert "【標準答案】" in teacher_html
    assert "【試題詳解】" in teacher_html


def test_clean_latex_to_unicode_and_deck_docx():
    from app.services import clean_latex_to_unicode, make_deck_docx
    from app.models import Deck, Slide

    # 1. Test math unicode conversion
    res_vec = clean_latex_to_unicode(r"$\vec{F} = m\vec{a}$")
    assert "F⃗" in res_vec and "a⃗" in res_vec

    res_grav = clean_latex_to_unicode(r"$F = G\frac{m_1 m_2}{r^2}$")
    assert "m₁" in res_grav and "r²" in res_grav

    res_energy = clean_latex_to_unicode(r"$E = mc^2 + \frac{1}{2}mv^2$")
    assert "c²" in res_energy and "v²" in res_energy

    # 2. Test deck docx generation
    test_deck = Deck(
        id="deck_docx_test",
        document_id="doc_1",
        title="牛頓運動定律與古典力學",
        subtitle="高中生｜45 分鐘",
        duration=45,
        mode="gemini",
        slides=[
            Slide(
                title="牛頓第二運動定律",
                bullets=["**公式**：$\\vec{F} = m\\vec{a}$", "質量為慣性的量度"],
                speaker_notes="各位同學好，今天我們要推導 $\\vec{F} = m\\vec{a}$ 的物理意義與應用。",
                visual_prompt="展示加速度與施力方向一致的示意圖",
                icon="🚀",
                source_pages=[5, 6],
            )
        ],
    )

    docx_bytes = make_deck_docx(test_deck)
    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 1000


def test_make_deck_slides_html():
    from app.services import make_deck_slides_html
    from app.models import Deck, Slide

    test_deck = Deck(
        id="deck_print_test",
        document_id="doc_1",
        title="國中數學：乘法分配律",
        subtitle="七年級上學期",
        duration=45,
        mode="gemini",
        slides=[
            Slide(
                title="乘法分配律核心",
                bullets=["**運算規則**：$a(b+c) = ab+ac$", "幾何意義：長方形面積分割"],
                speaker_notes="引導學生觀察長方形面積的拆解方式。",
                icon="🧮",
                source_pages=[12],
                visual_diagram={
                    "diagram_type": "key_formula",
                    "steps": [
                        {"label": "① 展開型態", "text": "$a(b+c) = ab + ac$"},
                        {"label": "② 幾何分割", "text": "總面積等於兩個小長方形面積之和"},
                        {"label": "③ 逆向因式分解", "text": "提取公因數 $ab + ac = a(b+c)$"},
                    ],
                    "takeaway": "乘法分配律是多項式運算與因式分解的核心基石。",
                },
            )
        ],
    )

    html = make_deck_slides_html(test_deck)
    assert "<!DOCTYPE html>" in html
    assert "國中數學：乘法分配律" in html
    assert "print-control-bar" in html
    assert "setOrientation('landscape')" in html
    assert "setOrientation('portrait')" in html
    assert "setLayout(this.value)" in html
    assert "toggleNotes(this.checked)" in html
    assert "window.print()" in html
    assert "katex.min.js" in html
    assert "renderMathInElement" in html
    assert "乘法分配律核心" in html
    assert "print-step-box" in html
    assert "① 展開型態" in html
    assert "核心結論與關鍵理解" in html
    assert "乘法分配律是多項式運算與因式分解的核心基石。" in html


def test_sanitize_latex_escapes_and_json():
    from app.services import sanitize_latex_escapes, _clean_and_load_json, _format_handout_text

    # 1. Test \x0c formfeed restoration
    corrupted_str = " \x0crac{c}{d} \\sim \\frac{c'}{d'}"
    repaired_str = sanitize_latex_escapes(corrupted_str)
    assert "\\frac{c}{d}" in repaired_str
    assert "\\sim" in repaired_str

    # 2. Test JSON parsing with bare LaTeX backslashes (\frac, \beta, \theta, \sim)
    raw_json = r'{"question": "已知 $\frac{c}{d} \sim \frac{c\'}{d\'}$，且 $\beta = \theta \times \nabla$", "answer": "B"}'
    parsed = _clean_and_load_json(raw_json)
    assert "\\frac{c}{d}" in parsed["question"]
    assert "\\beta" in parsed["question"]
    assert "\\theta" in parsed["question"]
    assert "\\times" in parsed["question"]
    assert "\\nabla" in parsed["question"]

    # 3. Test HTML formatter with bare and wrapped latex
    fmt = _format_handout_text("已知 \x0crac{c}{d} \\sim \\frac{c'}{d'}")
    assert "\\frac{c}{d}" in fmt

    # 4. Test complex bare formulas from user PDF
    fmt_fracab = _format_handout_text("分數 fracab 與 fraccd (b, d neq0)")
    assert "\\frac{a}{b}" in fmt_fracab
    assert "\\frac{c}{d}" in fmt_fracab
    assert "\\neq" in fmt_fracab

    fmt_left = _format_handout_text(r"加法定義為 \left[\frac{a}{b}\right] + \left[\frac{c}{d}\right] = \left[\frac{ad+bc}{bd}\right]")
    assert r"\left[\frac{a}{b}\right]" in fmt_left

    fmt_indices = _format_handout_text("指數映射 10^{-a} \\cdot 10^{-b} = 10^{-(a+b+1)} 及 2^{-2} times 5^{-2}")
    assert "10^{-a}" in fmt_indices
    assert "\\times" in fmt_indices

    fmt_series = _format_handout_text("迷思 frac13 + frac12 = frac25 及 frac78 與 frac38")
    assert "\\frac{1}{3}" in fmt_series
    assert "\\frac{1}{2}" in fmt_series
    assert "\\frac{2}{5}" in fmt_series
    assert "\\frac{7}{8}" in fmt_series
    assert "\\frac{3}{8}" in fmt_series


def test_auto_repair_math_expressions_times_and_symbols():
    from app.services import auto_repair_math_expressions, clean_latex_to_unicode

    # 1. Test \times and bare times operator
    res1 = auto_repair_math_expressions(r"2 \times 10^{-2}")
    assert res1 == r"$2 \times 10^{-2}$"

    res2 = auto_repair_math_expressions(r"2 times 10^{-2}")
    assert res2 == r"$2 \times 10^{-2}$"

    res3 = auto_repair_math_expressions(r"計算 3 times 5 的值")
    assert "$3 \\times 5$" in res3

    # 2. Test similarity symbol and prime fractions
    res4 = auto_repair_math_expressions(r"\frac{c}{d} \sim \frac{c'}{d'}")
    assert res4 == r"$\frac{c}{d} \sim \frac{c'}{d'}$"

    # 3. Test clean_latex_to_unicode for times and symbols
    uni1 = clean_latex_to_unicode(r"2 \times 10^{-2}")
    assert "2 × 10⁻²" == uni1

    uni2 = clean_latex_to_unicode(r"2 times 10^{-2}")
    assert "2 × 10⁻²" == uni2

    uni3 = clean_latex_to_unicode(r"\frac{c}{d} \sim \frac{c'}{d'}")
    assert "c/d ∼ c'/d'" == uni3

    # 4. Test \bar{x} and ASCII control character recovery
    res5 = auto_repair_math_expressions("獲得數據平均值為 \x08ar{x} = 4.05 \\times 10^{-6} m ，標準差為 \x08ar x 及 (\\frac{\\sigma}{\x08ar{x}})")
    assert r"\bar{x}" in res5
    assert r"\bar x" in res5
    assert r"\bar{x}" in res5

    # 5. Test JSON pre-escape & sanitization
    from app.services import _clean_and_load_json
    raw_json = '{"question": "平均值為 \\bar{x} = 4.05 \\times 10^{-6} m", "opt": "\\frac{\\sigma}{\\bar{x}}"}'
    parsed = _clean_and_load_json(raw_json)
    assert "\\bar{x}" in parsed["question"]
    assert "\\bar{x}" in parsed["opt"]


def test_make_deck_slides_html_print_and_katex():
    from app.models import Deck, Slide
    from app.services import make_deck_slides_html

    deck = Deck(
        id="d_print_test",
        document_id="doc_1",
        title="牛頓運動定律簡報",
        subtitle="高中物理",
        duration=45,
        mode="gemini",
        slides=[
            Slide(
                title="牛頓第二定律 $\\vec{F}=m\\vec{a}$",
                bullets=["合力與加速度成正比", "平均測量值 $\\bar{x} = 4.05 \\times 10^{-6}$"],
                speaker_notes="解說 $\\vec{F}=m\\vec{a}$ 與 $\\bar{x}$",
                visual_description="力與加速度向量示意圖",
                icon="🚀",
                source_pages=[1, 2],
            )
        ],
    )

    html_out = make_deck_slides_html(deck)
    assert "triggerPrint()" in html_out
    assert "katexOptions" in html_out
    assert "unescapeMathInElement" in html_out
    assert "document.fonts.ready" in html_out
    assert "牛頓第二定律" in html_out
    assert "A4 landscape" in html_out















